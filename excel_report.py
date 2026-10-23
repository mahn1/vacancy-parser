"""Сборка Excel-отчёта: лист с вакансиями и лист со сводкой."""

from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from transform import EXPERIENCE_ORDER, Vacancy, experience_label, salary_stats

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)
SECTION_FONT = Font(bold=True, size=12, color="1F4E78")
LINK_FONT = Font(color="0563C1", underline="single")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MONEY = "#,##0"

# (заголовок, ширина колонки)
COLUMNS = [
    ("Вакансия", 45),
    ("Компания", 35),
    ("Регион", 25),
    ("Зарплата от, ₽", 15),
    ("Зарплата до, ₽", 15),
    ("Опыт", 12),
    ("График", 28),
    ("Навыки", 50),
    ("Опубликована", 14),
    ("Ссылка", 12),
]


def _style_header(ws: Worksheet, row: int, columns: int) -> None:
    for col in range(1, columns + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def _write_vacancies(ws: Worksheet, vacancies: list[Vacancy]) -> None:
    ws.title = "Вакансии"
    for col, (title, width) in enumerate(COLUMNS, start=1):
        ws.cell(row=1, column=col, value=title)
        ws.column_dimensions[get_column_letter(col)].width = width
    _style_header(ws, 1, len(COLUMNS))
    ws.row_dimensions[1].height = 30

    for row, v in enumerate(vacancies, start=2):
        values = [
            v.title,
            v.company,
            v.region,
            v.salary_from,
            v.salary_to,
            experience_label(v.experience_years),
            v.schedule,
            ", ".join(v.skills),
            v.published,
            "Открыть" if v.url else None,
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=col in (1, 2, 8))

        ws.cell(row=row, column=4).number_format = MONEY
        ws.cell(row=row, column=5).number_format = MONEY
        ws.cell(row=row, column=9).number_format = "DD.MM.YYYY"
        if v.url:
            link = ws.cell(row=row, column=10)
            link.hyperlink = v.url
            link.font = LINK_FONT

    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(len(vacancies) + 1, 2)}"


def _write_table(
    ws: Worksheet,
    start_row: int,
    title: str,
    headers: list[str],
    rows: list[list],
    money_cols: tuple[int, ...] = (),
) -> int:
    """Пишет небольшую таблицу с заголовком раздела. Возвращает следующую свободную строку."""
    ws.cell(row=start_row, column=1, value=title).font = SECTION_FONT
    header_row = start_row + 1
    for col, header in enumerate(headers, start=1):
        ws.cell(row=header_row, column=col, value=header)
    _style_header(ws, header_row, len(headers))

    for r, values in enumerate(rows, start=header_row + 1):
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.border = BORDER
            if col in money_cols:
                cell.number_format = MONEY
    return header_row + len(rows) + 2


def _write_summary(ws: Worksheet, vacancies: list[Vacancy], query: str, region: str) -> None:
    ws.title = "Сводка"
    ws.column_dimensions["A"].width = 45
    for col in "BCD":
        ws.column_dimensions[col].width = 20

    ws["A1"] = f"Отчёт по вакансиям: «{query}»"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Регион: {region} · Источник: trudvsem.ru · Сформирован {datetime.now():%d.%m.%Y %H:%M}"
    ws["A2"].font = Font(italic=True, color="595959")

    stats = salary_stats(vacancies)
    row = _write_table(ws, 4, "Общие показатели", ["Показатель", "Значение"], [
        ["Всего вакансий", len(vacancies)],
        ["С указанной зарплатой", stats["count"]],
        ["Медианная зарплата, ₽", round(stats["median"]) if stats["median"] else "—"],
        ["Средняя зарплата, ₽", round(stats["mean"]) if stats["mean"] else "—"],
    ])
    ws["B8"].number_format = ws["B9"].number_format = MONEY

    by_experience = {label: [] for label in EXPERIENCE_ORDER}
    for v in vacancies:
        by_experience[experience_label(v.experience_years)].append(v)
    experience_rows = []
    for label, group in by_experience.items():
        if group:
            med = salary_stats(group)["median"]
            experience_rows.append([label, len(group), round(med) if med else None])
    exp_header = row + 1
    row = _write_table(
        ws, row, "Зарплата по опыту", ["Опыт", "Вакансий", "Медиана, ₽"], experience_rows, money_cols=(3,)
    )

    if experience_rows:
        chart = BarChart()
        chart.title = "Медианная зарплата по опыту, ₽"
        chart.style = 10
        chart.legend = None
        chart.height, chart.width = 7.5, 16
        values = Reference(ws, min_col=3, min_row=exp_header, max_row=exp_header + len(experience_rows))
        labels = Reference(ws, min_col=1, min_row=exp_header + 1, max_row=exp_header + len(experience_rows))
        chart.add_data(values, titles_from_data=True)
        chart.set_categories(labels)
        ws.add_chart(chart, "F4")

    employers = Counter(v.company for v in vacancies).most_common(10)
    row = _write_table(ws, row, "Топ-10 работодателей", ["Компания", "Вакансий"], [list(e) for e in employers])

    skills = Counter(skill for v in vacancies for skill in v.skills).most_common(15)
    if skills:
        _write_table(ws, row, "Топ-15 навыков", ["Навык", "Упоминаний"], [list(s) for s in skills])


def build_report(vacancies: list[Vacancy], path: Path, query: str, region: str) -> Path:
    wb = Workbook()
    _write_summary(wb.active, vacancies, query, region)
    _write_vacancies(wb.create_sheet(), vacancies)
    wb.save(path)
    return path
